from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "FLOW_provozni_prirucka.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(89, 99, 112)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_fixed_widths(table, widths_inches):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = table_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        table_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    for row in table.rows:
        for cell, width in zip(row.cells, widths_inches):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(round(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_font(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def add_body(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(18 if level == 1 else 14)
    p.paragraph_format.space_after = Pt(10 if level == 1 else 7)
    run = p.add_run(text)
    set_font(run, size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_fixed_widths(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(f"{label}: "), size=10.5, color=DARK_BLUE, bold=True)
    set_font(p.add_run(text), size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, color in ((1, 16, BLUE), (2, 13, BLUE), (3, 12, DARK_BLUE)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("FLOW | provozní příručka"), size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Interní pracovní podklad | verze 1.0 | 11. 8. 2026"), size=8.5, color=MUTED)


def build_document():
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(22)
    title.paragraph_format.space_after = Pt(6)
    set_font(title.add_run("FLOW"), size=30, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    set_font(subtitle.add_run("Provozní příručka pro záznam, triáž a předávku směny"), size=14, color=MUTED)

    meta = doc.add_table(rows=1, cols=3)
    set_fixed_widths(meta, [2.1, 2.2, 2.2])
    for cell, (label, value) in zip(meta.rows[0].cells, (("Účel", "zachytit znalost"), ("Rytmus", "průběžně + předávka"), ("Vlastník", "vedoucí úseku"))):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(label + "\n"), size=9, color=MUTED, bold=True)
        set_font(p.add_run(value), size=10.5, color=INK, bold=True)

    add_heading(doc, "Jak FLOW používat", 1)
    add_body(doc, "FLOW slouží k rychlému zachycení zkušeností přímo z výroby. Každý záznam má být krátký, dohledatelný a při problému dostatečně konkrétní, aby na něj mohla navázat další směna.")
    for item in (
        "Zachytit: zvolit typ záznamu, napsat větu a případně přidat fotku.",
        "Upřesnit: doplnit TL, úsek, stroj, prioritu, parametr a štítky.",
        "Třídit: v Hledat použít rychlé filtry Nevyřešené, Kritické, S fotkou nebo Poslední týden.",
        "Vyřešit: u problému doplnit řešení, označit ho jako ověřené a problém uzavřít až po ověření v provozu.",
        "Předat: před koncem směny projít pracovní frontu a otevřené kritické záznamy.",
    ):
        add_number(doc, item)

    add_callout(doc, "Pravidlo kvality", "Zápis má říkat co se stalo, kde, u čeho a jaký byl dopad. Fotka nenahrazuje krátký popis.")

    add_heading(doc, "Triáž priorit", 1)
    priority_table = doc.add_table(rows=1, cols=3)
    priority_table.style = "Table Grid"
    set_fixed_widths(priority_table, [1.35, 2.2, 2.95])
    for cell, text in zip(priority_table.rows[0].cells, ("Priorita", "Kdy ji použít", "Další krok")):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        set_font(p.add_run(text), size=10, color=DARK_BLUE, bold=True)
    for row in (
        ("Kritická", "Bezpečnost, odstávka, riziko zmetku", "Okamžitě otevřít pracovní frontu a předat odpovědné osobě."),
        ("Důležitá", "Opakující se odchylka nebo dopad na kvalitu", "Doplnit kontext, termín ověření a řešení."),
        ("Běžná", "Zlepšení, poznatek, doporučené nastavení", "Zařadit k TL nebo stroji pro budoucí dohledání."),
    ):
        cells = priority_table.add_row().cells
        for cell, text in zip(cells, row):
            p = cell.paragraphs[0]
            set_font(p.add_run(text), size=10.2, color=INK)

    add_heading(doc, "Předávka směny", 1)
    add_body(doc, "Předávka má trvat několik minut a vychází z domovské pracovní fronty. Cílem není opisovat vše, ale zajistit, že podstatný problém nezůstane bez vlastníka.")
    checklist = doc.add_table(rows=1, cols=2)
    checklist.style = "Table Grid"
    set_fixed_widths(checklist, [0.65, 5.85])
    for cell, text in zip(checklist.rows[0].cells, ("✓", "Kontrolní bod")):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if text == "✓" else WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(text), size=10, color=DARK_BLUE, bold=True)
    for item in (
        "Otevřené kritické problémy mají jasný stav a další krok.",
        "Nové problémy mají přiřazené TL, úsek nebo stroj, pokud je známý.",
        "Použitelná řešení jsou označena jako ověřená.",
        "Záznamy pro další směnu obsahují krátký kontext, ne jen interní zkratku.",
    ):
        cells = checklist.add_row().cells
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cells[0].paragraphs[0].add_run("☐"), size=11, color=DARK_BLUE)
        set_font(cells[1].paragraphs[0].add_run(item), size=10.2, color=INK)

    add_heading(doc, "Data a obnova", 1)
    add_body(doc, "Aplikace funguje i bez připojení. Pro ochranu záznamů používejte obrazovku Data: export JSON je úplná záloha včetně fotografií; CSV a XLSX slouží hlavně pro sdílení a analýzu. Obnovu JSON provádějte jen ze známé zálohy, protože nahrazuje současná data v aplikaci.")
    add_callout(doc, "Doporučený rytmus", "Úplná JSON záloha jednou týdně a před každým větším testem nebo hromadnou úpravou záznamů.")

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
